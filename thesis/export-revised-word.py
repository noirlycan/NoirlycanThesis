#!/usr/bin/env python3
"""Export the edited thesis sources to an editable UET-oriented DOCX."""

from __future__ import annotations

import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "luan-van-bien-tap.docx"


def clean(text: str) -> str:
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\{width=[^}]+\}", "", text)
    return (text.replace("✅", "Có").replace("❌", "Không")
            .replace("⚠️", "Cần kiểm tra").replace("⏳", "Chưa hoàn tất")
            .replace("✓", "Có"))


def add_page_number(section) -> None:
    paragraph = section.header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, end):
        paragraph.add_run()._r.append(element)


def add_toc(document: Document) -> None:
    document.add_heading("MỤC LỤC", level=1)
    paragraph = document.add_paragraph()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate):
        paragraph.add_run()._r.append(element)
    paragraph.add_run("Cập nhật mục lục trong Word trước khi nộp.")
    paragraph.add_run()._r.append(end)


def add_cover(document: Document) -> None:
    rows = [
        ("ĐẠI HỌC QUỐC GIA HÀ NỘI", 14, True, 0),
        ("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ", 14, True, 0),
        ("PHAN VĂN CƠ", 14, True, 50),
        ("TĂNG CƯỜNG PHÁT HIỆN MÃ ĐỘC TRƯỚC CÁC KĨ THUẬT NÉ TRÁNH", 16, True, 70),
        ("LUẬN VĂN THẠC SĨ", 14, True, 60),
        ("Ngành/chuyên ngành: [CẦN XÁC NHẬN THEO QUYẾT ĐỊNH ĐÀO TẠO]", 13, False, 30),
        ("Người hướng dẫn khoa học: [CẦN ĐIỀN]", 13, False, 0),
        ("HÀ NỘI – 2026", 14, True, 65),
    ]
    for value, size, bold, before in rows:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(before)
        run = paragraph.add_run(value)
        run.bold = bold
        run.font.size = Pt(size)
    document.add_page_break()


def render_markdown(document: Document, source: str) -> None:
    lines = source.splitlines()
    paragraph: list[str] = []
    code: list[str] = []
    in_code = False
    in_openxml = False

    def flush() -> None:
        if paragraph:
            value = clean(" ".join(paragraph))
            p = document.add_paragraph(value)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.clear()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        value = line.strip()
        if value.startswith("```{=openxml}"):
            flush()
            in_openxml = True
            i += 1
            continue
        if in_openxml:
            if value.startswith("```"):
                in_openxml = False
            i += 1
            continue
        if value.startswith("<!--"):
            i += 1
            continue
        if value.startswith("```"):
            flush()
            if in_code:
                p = document.add_paragraph()
                p.style = "No Spacing"
                run = p.add_run("\n".join(code))
                run.font.name = "Courier New"
                run.font.size = Pt(8)
                code.clear()
            in_code = not in_code
            i += 1
            continue
        if in_code:
            code.append(line)
            i += 1
            continue
        if not value or value == "---":
            flush()
            i += 1
            continue
        match = re.match(r"^(#{1,4})\s+(.+)$", value)
        if match:
            flush()
            document.add_heading(clean(match.group(2)), level=len(match.group(1)))
            i += 1
            continue
        figure = re.match(r"^!\[([^]]*)\]\(([^)]+)\)(?:\{[^}]*\})?$", value)
        if figure:
            flush()
            image = (ROOT / figure.group(2)).resolve()
            if image.is_file() and ROOT in image.parents:
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(image), width=Cm(14))
            caption = document.add_paragraph(clean(figure.group(1)))
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.style = "Caption"
            i += 1
            continue
        if value.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*[-:]+\s*\|", lines[i + 1].strip()):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [clean(part.strip()) for part in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+", part) for part in cells):
                    rows.append(cells)
                i += 1
            if rows:
                width = max(len(row) for row in rows)
                table = document.add_table(rows=len(rows), cols=width)
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r, cells in enumerate(rows):
                    for c, cell in enumerate(cells):
                        table.cell(r, c).text = cell
                        for p in table.cell(r, c).paragraphs:
                            p.paragraph_format.space_after = Pt(0)
                            for run in p.runs:
                                run.font.size = Pt(9)
                                if r == 0:
                                    run.bold = True
            continue
        item = re.match(r"^(?:[-*]|\d+\.)\s+(.+)$", value)
        if item:
            flush()
            p = document.add_paragraph(clean(item.group(1)), style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            i += 1
            continue
        if value.startswith("$$"):
            flush()
            document.add_paragraph(clean(value))
            i += 1
            continue
        if value.startswith("**Bảng") or value.startswith("**Hình"):
            flush()
            p = document.add_paragraph(clean(value))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if value.startswith("**Hình") else WD_ALIGN_PARAGRAPH.LEFT
            p.style = "Caption"
            i += 1
            continue
        paragraph.append(value)
        i += 1
    flush()


def main() -> None:
    subprocess.run(["python3", str(ROOT / "assemble-htw.py")], check=True, cwd=ROOT)
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.header_distance = Cm(1)
    add_page_number(section)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(13)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(7)
    for name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = None
        style.font.size = Pt(14 if name == "Heading 1" else 13)
    document.styles["Heading 1"].paragraph_format.page_break_before = True
    document.styles["Heading 1"].paragraph_format.keep_with_next = True
    add_cover(document)
    declaration = (ROOT / "00-front-matter.md").read_text(encoding="utf-8")
    render_markdown(document, declaration[declaration.index("# LỜI CAM ĐOAN"):])
    add_toc(document)
    render_markdown(document, (ROOT / "02-danh-muc-viet-tat.md").read_text(encoding="utf-8"))
    sources = [
        "chapters/01-mo-dau.md", "chapters/02-co-so-ly-thuyet.md",
        "chapters/_assembled/03-phuong-phap-htw.md", "chapters/_assembled/04-thiet-ke-htw.md",
        "chapters/_assembled/05-thuc-nghiem-htw.md", "ket-luan.md",
        "danh-muc-cong-trinh.md", "tai-lieu-tham-khao.md", "phu-luc.md",
    ]
    for source in sources:
        render_markdown(document, (ROOT / source).read_text(encoding="utf-8"))
    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
